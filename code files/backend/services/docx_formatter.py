from __future__ import annotations

import io
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt

from ..utils.text import markdown_to_runs, sanitize_text, terms_to_list

PROJECT_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_LOGO = PROJECT_ROOT / "assets" / "logo.png"


def _set_run_font(run, name: str = "Times New Roman", size: float = 11, bold: bool = False, italic: bool = False) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def _add_markdown_runs(paragraph, text: str, size: float = 11, bold: bool = False) -> None:
    for chunk, is_bold, is_italic, is_code in markdown_to_runs(text):
        run = paragraph.add_run(chunk)
        _set_run_font(run, size=size, bold=bold or is_bold, italic=is_italic)
        if is_code:
            run.font.name = "Courier New"


def _set_cell(cell, value: str, bold: bool = False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(sanitize_text(value))
    _set_run_font(r, size=10, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def format_docx(text: str, doc_type: str, terms: str = "", logo_path: str | Path | None = None) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    logo = Path(logo_path) if logo_path else DEFAULT_LOGO
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if logo.exists():
        header.add_run().add_picture(str(logo), width=Inches(2.35))
    else:
        r = header.add_run("LegalEase")
        _set_run_font(r, size=12, bold=True)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer.add_run("LegalEase - AI-powered legal document draft | Review before use")
    _set_run_font(fr, size=8)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(4)
    title.paragraph_format.space_after = Pt(12)
    _add_markdown_runs(title, doc_type, size=16, bold=True)

    for raw_line in sanitize_text(text).splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if line.startswith("# "):
            continue
        if line.startswith("### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            _add_markdown_runs(p, line[4:], size=12, bold=True)
            continue
        if line.startswith("## "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            _add_markdown_runs(p, line[3:], size=13, bold=True)
            continue
        if re.match(r"^[-*] ", line):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent = Inches(0.25)
            _add_markdown_runs(p, line[2:])
            continue
        if re.match(r"^\d+[.)] ", line):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.15)
            _add_markdown_runs(p, line)
            continue
        p = doc.add_paragraph()
        _add_markdown_runs(p, line)

    cleaned_terms = terms_to_list(terms)
    if cleaned_terms:
        doc.add_paragraph()
        p = doc.add_paragraph()
        _add_markdown_runs(p, "Terms Summary", size=12, bold=True)
        table = doc.add_table(rows=1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = "Table Grid"
        _set_cell(table.rows[0].cells[0], "#", bold=True)
        _set_cell(table.rows[0].cells[1], "Term", bold=True)
        for idx, term in enumerate(cleaned_terms, 1):
            cells = table.add_row().cells
            _set_cell(cells[0], str(idx))
            _set_cell(cells[1], term)

    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()
