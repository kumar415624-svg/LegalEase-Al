from io import BytesIO

from docx import Document

from backend.services.exporters import format_docx, format_pdf, format_txt

SAMPLE = """# Non-Disclosure Agreement

## 1. Confidentiality
- Confidential information must be protected.
- Return materials within 30 days.

## 2. Signatures
Party 1: ______________________________
"""


def test_txt_export():
    data = format_txt(SAMPLE)
    assert data.startswith(b"Non-Disclosure Agreement")


def test_docx_export_contains_content_and_terms_table():
    data = format_docx(SAMPLE, "Non-Disclosure Agreement", "Protect information; Return materials")
    assert data[:2] == b"PK"
    doc = Document(BytesIO(data))
    paragraph_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Confidentiality" in paragraph_text
    assert len(doc.tables) == 1
    assert "Terms Summary" in paragraph_text


def test_pdf_export_is_valid_header():
    data = format_pdf(SAMPLE, "Non-Disclosure Agreement", "Protect information; Return materials")
    assert data.startswith(b"%PDF")


def test_docx_export_applies_bold_markdown():
    data = format_docx("## Terms\nThis is **very important**.", "Agreement")
    doc = Document(BytesIO(data))
    runs = [run for p in doc.paragraphs for run in p.runs if "very important" in run.text]
    assert runs
    assert runs[0].bold is True


def test_txt_export_strips_markdown_markers():
    data = format_txt("# Agreement\nThis is **bold**.")
    text = data.decode("utf-8")
    assert "**" not in text
    assert "Agreement" in text


def test_pdf_export_handles_unicode_and_markdown_without_crashing():
    data = format_pdf("# Agreement 🚀\n\nThis is **bold** and café — café.", "Agreement 🚀")
    assert data.startswith(b"%PDF")
